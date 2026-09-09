from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "deliverables"
OUT_DIR.mkdir(exist_ok=True)
OUT = OUT_DIR / "水下六足机器人项目协作仓库说明与导师汇报指南.docx"


def set_run_font(run, east_asia="Microsoft YaHei", western="Arial", size=None, bold=None, color="000000"):
    run.font.name = western
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), western)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), western)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color="666666")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    r = paragraph.add_run()._r
    r.append(fld_char1)
    r.append(instr)
    r.append(fld_char2)
    run2 = paragraph.add_run(" 页")
    set_run_font(run2, size=9, color="666666")


def add_para(doc, text="", style=None, bold_lead=None, align=None, keep=False):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    if keep:
        set_keep_with_next(p)
    return p


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_numbered_list(doc, items):
    for index, text in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.28)
        r = p.add_run(f"{index}.  {text}")
        set_run_font(r)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0]
    repeat_table_header(hdr)
    for i, value in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = ""
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "1F4E78")
        set_cell_margins(cell)
        set_cell_borders(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(value)
        set_run_font(r, size=9.5, bold=True, color="FFFFFF")
        if widths:
            cell.width = Inches(widths[i])
    for row_index, row_values in enumerate(rows):
        row = table.add_row()
        for i, value in enumerate(row_values):
            cell = row.cells[i]
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_borders(cell)
            if row_index % 2 == 1:
                set_cell_shading(cell, "F3F6F9")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if i == 0 and len(headers) <= 4:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(value))
            set_run_font(r, size=9.2)
            if widths:
                cell.width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(10.8)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.line_spacing = 1.28
normal.paragraph_format.space_after = Pt(6)

for style_name, size, before, after in (
    ("Title", 24, 0, 14),
    ("Heading 1", 16, 14, 7),
    ("Heading 2", 12.5, 10, 5),
    ("Heading 3", 11.3, 8, 4),
):
    s = styles[style_name]
    s.font.name = "Arial"
    s._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True

title_ppr = styles["Title"]._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

for sec in doc.sections:
    add_page_number(sec.footer.paragraphs[0])

# Cover page
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(82)
r = p.add_run("水下六足机器人项目协作仓库说明与导师汇报指南")
set_run_font(r, size=24, bold=True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(8)
r = p.add_run("用于说明当前成果 组织两地协作 明确下一阶段工作")
set_run_font(r, size=12, color="444444")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(54)
for label in ("汇报人  ____________________", "指导教师  ____________________", "日期  2026 年 9 月 8 日"):
    r = p.add_run(label + "\n")
    set_run_font(r, size=11)

doc.add_paragraph().paragraph_format.space_before = Pt(55)
add_para(
    doc,
    "本次工作的成果是一套水下六足机器人项目的协作仓库框架。它已经明确两地团队如何交付机械参数、如何对接控制程序、如何记录任务，以及每个阶段怎样验收。目前尚未完成可运行的 MuJoCo 机器人模型和控制代码。下一步需要导师确认项目范围、双方分工、关键机械方案和第一阶段交付时间。",
)

doc.add_page_break()

doc.add_heading("先向导师说明什么", level=1)
add_para(
    doc,
    "向导师汇报时，先说结论：我已经把项目协作所需的目录、接口文件、任务看板和验收标准整理成了一个仓库框架。这个框架解决的是两地团队怎样交换信息和共同开发的问题，不代表机器人仿真已经完成。",
)
add_para(doc, "导师最需要知道的四件事", style="Heading 2", keep=True)
for item in (
    "目前完成的是项目组织与接口设计，属于启动阶段成果。",
    "机械尺寸、质量、惯量、关节范围、折叠结构和吸附参数仍需机电团队提供。",
    "智能团队将在参数确认后建立 MuJoCo 模型，并实现可视化和基础关节控制。",
    "希望导师确认第一阶段范围、负责人和验收时间，避免两地团队重复修改。",
):
    add_bullet(doc, item)

doc.add_heading("已经完成和尚未完成的内容", level=1)
add_table(
    doc,
    ["类别", "当前状态", "准确说法"],
    [
        ("协作仓库目录", "已完成", "已建立项目文件夹、配置目录、模型目录、代码目录和测试目录"),
        ("机械控制接口", "初稿已完成", "已列出双方必须交换的参数，等待实际数据填入"),
        ("任务与验收流程", "初稿已完成", "已建立 M0 到 M4 任务和验收标准，等待导师确认"),
        ("控制序列格式", "模板已完成", "已设计地面到墙面切换的状态和参数字段，数值仍未确定"),
        ("MuJoCo 机器人模型", "尚未完成", "需要机械参数和结构方案后才能准确建模"),
        ("可视化与控制代码", "尚未完成", "将在模型合同确认后开始实现"),
        ("平面和墙面实验", "尚未开始", "属于后续里程碑，不能作为当前成果汇报"),
    ],
    widths=[1.3, 1.25, 4.05],
)

doc.add_heading("为什么先做这套仓库", level=1)
add_para(
    doc,
    "两地团队最常见的问题不是不会建模，而是使用了不同的尺寸、坐标系、关节名称和版本。机械模型一旦改变，控制程序可能全部返工。因此第一步应把输入、输出和验收方法写清楚，让每次修改都有文件和版本可查。",
)

doc.add_heading("各文件的解释和作用", level=1)
add_para(doc, "导师不需要逐行阅读所有文件。可以先打开 Word 指南，再看 README 和两个 YAML 配置文件。")
add_table(
    doc,
    ["文件", "通俗解释", "主要使用者", "当前作用"],
    [
        ("README.md", "项目总说明书和入口", "所有成员", "说明项目目标、目录、分工和启动顺序"),
        ("PROJECT_BOARD.md", "任务看板", "两地接口人", "记录每项任务由谁负责、交付什么、谁验收以及是否阻塞"),
        ("docs/01_project_charter.md", "项目章程", "导师和全体成员", "限定当前阶段范围，并划分机电团队和智能团队的责任"),
        ("docs/02_interface_contract.md", "机械与控制交接说明", "机械建模和控制成员", "统一坐标系、关节命名、单位以及双方必须提供的参数"),
        ("docs/03_milestones_acceptance.md", "分阶段验收表", "导师和验收人", "规定模型、行走、墙面运动和切换分别达到什么条件才算完成"),
        ("docs/04_remote_workflow.md", "异地协作规则", "两地接口人", "规定开会频率、交付规则、版本管理和问题升级方式"),
        ("config/model_contract.yaml", "可由程序读取的机械参数表", "机电团队填写 智能团队使用", "保存尺寸、质量、惯量、关节范围、折叠结构和吸附参数"),
        ("config/control_sequence_template.yaml", "动作流程参数表", "控制成员", "保存平面到墙面切换的各个状态、动作条件、超时和失败处理"),
        ("docs/decisions", "关键决定档案", "导师和接口人", "记录为什么选择某个坐标系、结构方案或开发范围"),
        ("assets models src tests runs", "模型资源和程序目录", "研发成员", "分别存放网格、MuJoCo 模型、控制代码、测试和运行结果"),
    ],
    widths=[1.55, 1.75, 1.45, 2.25],
)

doc.add_heading("几个基础概念", level=2)
for lead, explanation in (
    ("仓库", "可以理解为一个有版本记录的项目文件夹。成员可以知道谁在什么时候修改了什么。"),
    ("Markdown 文件", "扩展名是 md，内容主要给人阅读。用记事本、VS Code 或网页都可以打开。"),
    ("YAML 文件", "扩展名是 yaml，既方便人填写，也方便 Python 程序读取。缩进代表层级，不能随意删空格。"),
    ("模型合同", "不是法律合同，而是机械团队和控制团队共同遵守的数据格式。只要字段和单位一致，双方可以并行工作。"),
    ("TBD", "表示尚未确定。必须由相应负责人补充，不能用猜测值直接进入正式实验。"),
):
    add_para(doc, f"{lead}  {explanation}", bold_lead=lead)

doc.add_heading("五分钟汇报讲稿", level=1)
add_para(doc, "下面的内容可以直接照着讲，括号中的提示不用读出来。")

doc.add_heading("第一分钟介绍问题", level=2)
add_para(
    doc,
    "老师好，我们的项目是设计一款水下六足机器人。机器人中央平台可以折叠，六条腿的末端具有磁吸或粘附单元。最终目标是实现平面行走、墙面运动，以及从平面向墙面的自动切换。因为智能团队和机电团队在不同地点，目前最大的困难是机械模型和控制程序之间缺少统一的交接方式。",
)

doc.add_heading("第二分钟介绍本次成果", level=2)
add_para(
    doc,
    "我先整理了一套项目协作仓库框架。目前已经建立项目目录、双方分工、机械与控制接口、任务看板、阶段验收标准，以及平面到墙面切换的控制参数模板。这些文件的目的，是让两地成员以后通过同一套名称、单位和版本协作。",
)

doc.add_heading("第三分钟说明接口设计", level=2)
add_para(
    doc,
    "在接口文件中，我统一了 base_link 坐标系、六条腿和关节的命名方式，并列出了机电团队需要提供的尺寸、质量、惯量、关节范围、折叠机构和吸附参数。智能团队收到这些数据后，才能建立准确的 MuJoCo 模型并开发控制器。",
)

doc.add_heading("第四分钟说明技术路线", level=2)
add_para(
    doc,
    "第一阶段先不接 ROS，也暂不加入复杂水动力。我们先完成 MuJoCo 模型可视化、关节位置控制和默认站立姿态。之后再做平面步态、墙面吸附运动，最后用状态机完成平面到墙面的切换。切换过程中每个状态的目标关节值、吸附命令、结束条件和超时参数都会保存在 YAML 文件中，便于以后映射到实体机器人。",
)

doc.add_heading("第五分钟提出请求", level=2)
add_para(
    doc,
    "目前这些文件还是启动阶段初稿。我想请老师帮助确认四点：第一阶段是否按 MuJoCo 可视化和基础控制推进；双方的职责划分是否合适；机电团队由谁提供模型参数以及截止时间；两地是否每周安排一次固定的技术对接。确认后，我会根据实际机械参数开始搭建可运行模型。",
)

doc.add_heading("汇报时建议展示的顺序", level=1)
add_numbered_list(doc, (
    "先展示 README，让导师看见项目目标和第一周启动顺序。",
    "再展示 model_contract.yaml，指出这里不是程序，而是等待机电团队填写的统一参数表。",
    "接着展示 PROJECT_BOARD.md，说明每个阻塞任务都有负责人和交付物。",
    "最后展示 milestones_acceptance.md，说明以后将按可运行结果和日志验收。",
))

doc.add_heading("需要导师确认的问题", level=1)
add_para(doc, "汇报结束后，不要只问老师有没有意见。可以按下面的顺序请导师做决定。")
add_table(
    doc,
    ["优先级", "需要确认的事项", "为什么现在要确认", "建议结果"],
    [
        ("1", "第一阶段范围", "范围不清会同时开展水动力、ROS 和硬件通信，难以按时产出", "先完成 MuJoCo 可视化、基础关节控制和日志"),
        ("2", "机械方案负责人", "模型尺寸和折叠拓扑无人最终确认，智能团队无法准确建模", "指定一名机电学生接口人和一名教师决策人"),
        ("3", "关键参数截止时间", "所有仿真工作依赖模型合同", "确定 model_contract.yaml 的填写日期"),
        ("4", "足端技术路线", "磁吸和粘附对应不同墙面材料及力学模型", "确认首选方案或可替换模块方案"),
        ("5", "例会与验收频率", "两地信息只停留在聊天记录中容易遗漏", "每周一次技术对接 每两周一次教师决策"),
    ],
    widths=[0.65, 1.45, 2.6, 2.3],
)

doc.add_heading("导师可能会问的问题", level=1)
qa = [
    ("为什么现在没有机器人仿真画面", "因为准确的 MuJoCo 模型需要机械尺寸、质量、惯量、关节轴和折叠结构。目前先把参数交接格式固定下来，避免使用猜测模型后返工。"),
    ("这些文件是程序吗", "大部分是说明和配置文件。YAML 文件以后会被 Python 控制程序读取。models 和 src 目录将分别存放 MuJoCo 模型和控制代码。"),
    ("为什么不用 ROS", "项目要求当前阶段跳过 ROS。先在 MuJoCo 内部打通模型和控制链路，降低初期复杂度。"),
    ("吸附在 MuJoCo 中怎么做", "初期可以用约束或等效力近似，但需要先确定最大法向力、切向力、建立延迟、释放延迟和失效条件。后续再根据实验数据提高模型精度。"),
    ("怎样证明下一阶段完成", "模型必须能够无报错加载，关节名称和范围与参数合同一致，默认姿态能稳定保持，并能输出 base 位姿、关节状态、控制量和接触状态日志。"),
    ("你个人接下来做什么", "先根据导师意见修改接口文件，联系机电接口人补全模型合同，然后建立 MJCF 骨架和最小地面场景。"),
]
for q, a in qa:
    p = doc.add_paragraph()
    r = p.add_run("问  " + q)
    set_run_font(r, bold=True)
    p.paragraph_format.keep_with_next = True
    add_para(doc, "答  " + a)

doc.add_heading("如何把成果发给导师", level=1)
add_para(
    doc,
    "最简单的方式是发送本次生成的压缩包。导师解压后先打开这份 Word 指南，再根据需要查看 README 和配置文件。不要只发送十几个零散文件，也不要只发截图。",
)

doc.add_heading("使用微信或 QQ 发送", level=2)
add_numbered_list(doc, (
    "找到 deliverables 文件夹中的导师汇报包 zip 文件。",
    "直接把压缩包和一段简短说明发给导师。",
    "如果聊天软件限制文件大小，把压缩包上传到学校网盘或常用云盘，再发送链接。",
    "导师反馈后，把结论写入会议纪要或决策记录，不要让决定只留在聊天记录中。",
))

doc.add_heading("使用邮件发送", level=2)
add_numbered_list(doc, (
    "邮件主题写清项目和内容，例如 水下六足机器人项目协作框架初稿 请审阅。",
    "正文先说明当前完成了什么，再说明尚未完成什么。",
    "附件放压缩包。如果附件过大，改用网盘链接，并写明有效期。",
    "正文最后列出希望导师确认的具体事项。",
))

doc.add_heading("以后使用 Git 仓库协作", level=2)
add_para(
    doc,
    "压缩包适合第一次汇报，但多人长期修改时容易出现多个版本。导师确认方案后，可以把项目上传到学校 GitLab、Gitee 或 GitHub。两地团队都从同一仓库获取文件，每次修改都有记录。上传到外部平台前，应先确认项目是否涉及保密要求以及哪些 CAD 和实验数据允许公开。",
)

doc.add_heading("文件发送前检查", level=1)
for item in (
    "压缩包能够正常打开，里面包含 Word 指南、README、docs、config 和任务看板。",
    "文件名包含项目名称和日期，不使用 最终版 新版 等模糊名称。",
    "没有把个人密码、访问令牌、账号信息或不应公开的数据打包进去。",
    "聊天或邮件正文说明当前只是协作框架，不声称 MuJoCo 模型已经运行。",
    "明确写出希望导师确认的四项决定和期望反馈时间。",
):
    add_para(doc, "□  " + item)

doc.add_heading("可直接复制的消息", level=1)

doc.add_heading("微信简短版本", level=2)
add_para(
    doc,
    "老师您好，我根据水下六足机器人项目目前两地协作不便的问题，整理了一套项目协作仓库初稿，包括双方分工、机械与控制接口、任务看板、阶段验收标准，以及平面到墙面切换的控制参数模板。目前完成的是协作和接口框架，MuJoCo 模型及控制代码尚未开始正式实现。我把说明文档和全部文件打包发给您，想请您重点确认第一阶段范围、双方接口负责人、机械参数交付时间和例会安排。",
)

doc.add_heading("邮件正式版本", level=2)
add_para(doc, "邮件主题  水下六足机器人项目协作框架初稿 请审阅", bold_lead="邮件主题")
add_para(doc, "老师您好")
add_para(
    doc,
    "针对智能团队与机电团队异地协作、模型参数和任务交接不统一的问题，我整理了一套水下六足机器人项目协作仓库初稿。仓库目前包含项目章程、机械与控制接口协议、任务看板、阶段验收标准、异地协作流程，以及平面到墙面切换的控制参数模板。",
)
add_para(
    doc,
    "当前成果主要用于统一项目输入、输出和版本，尚未完成可运行的 MuJoCo 模型与控制程序。待机电团队补全尺寸、质量、惯量、关节和折叠机构等信息后，我们再开展模型可视化和基础控制。",
)
add_para(doc, "想请您帮助确认以下事项")
for item in (
    "第一阶段是否限定为 MuJoCo 模型可视化、基础关节控制和运行日志。",
    "双方学生接口人和教师决策人如何安排。",
    "机电团队补全模型参数的负责人和时间。",
    "两地团队是否建立每周一次的固定技术对接。",
):
    add_bullet(doc, item)
add_para(doc, "附件为项目汇报包，请您审阅。谢谢老师。")
add_para(doc, "学生  __________\n日期  __________")

doc.add_heading("会议结束后的反馈消息", level=2)
add_para(
    doc,
    "老师您好，我把今天会议确认的范围、负责人、截止时间和待解决问题记录到了项目仓库。接下来我会先完成接口文件修改，并配合机电同学补全模型参数。参数确认后开始建立 MuJoCo 模型。若我的记录有遗漏，请您指出。",
)

doc.add_heading("收到导师意见后怎么做", level=1)
add_numbered_list(doc, (
    "把导师确认的决定写入 docs/decisions，不依靠记忆保存。",
    "在 PROJECT_BOARD.md 中填写真实负责人、截止日期和任务状态。",
    "邀请机电接口人补全 config/model_contract.yaml，不确定的数值继续保留 TBD。",
    "完成参数评审后建立 models/robot.xml，并用最小场景检查模型能否加载。",
    "实现关节位置控制和日志记录，再按 M1 标准向导师演示。",
))

doc.add_heading("第一次正式对接建议", level=1)
add_para(doc, "第一次两地会议只解决接口，不讨论所有后续算法。会议应形成以下结果。")
for item in (
    "确认六足各有几个自由度，以及每个关节的轴和正方向。",
    "确认中央平台怎样折叠、转轴在哪里、最大折叠角是多少。",
    "确认足端首选磁吸还是粘附，以及墙面材料。",
    "确认机械网格、质量、惯量和关节限制的交付格式与日期。",
    "确定下一次验收内容为 MuJoCo 模型成功加载和关节名称检查。",
):
    add_bullet(doc, item)

doc.add_heading("当前汇报包内容", level=1)
add_para(
    doc,
    "汇报包包含本指南、项目入口、任务看板、项目章程、接口协议、验收标准、远程协作流程、模型参数合同、控制序列模板，以及为后续模型和代码预留的目录。它可以作为第一次导师汇报和两地启动会的材料。",
)

doc.core_properties.title = "水下六足机器人项目协作仓库说明与导师汇报指南"
doc.core_properties.subject = "水下六足机器人科研项目导师汇报和异地协作说明"
doc.core_properties.author = "河海大学水下六足机器人项目组"
doc.core_properties.keywords = "水下六足机器人 MuJoCo 协作仓库 导师汇报"

doc.save(OUT)
print(OUT)
